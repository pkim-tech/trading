"""Coverage for the report-provenance stamp added 2026-08-29, filed against
the real 2026-08-23 incident (docs/backlog_cache.md:150-151) -- a generated
report (output/full_review_gt_final_combined_20260823_202719.xlsx) was
trusted as current for over an hour after two commits changed the numbers it
contained, with nothing on the report itself flagging staleness. Stamp-only:
just proves the commit/dirty/timestamp line lands in both xlsx (Column
Definitions sheet) and csv (leading `#` comment) output, no comparison/alert
logic involved."""
import re
import sqlite3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from scripts.candidate_summary_report import (
    _write_csv, _write_xlsx, _git_provenance_stamp as _stamp_summary, COLUMN_DEFS,
)
from scripts.candidate_full_review import _git_provenance_stamp as _stamp_full_review

_COMMIT_RE = re.compile(r"^commit=([0-9a-f]{12}(\*dirty\*)?|unknown) at \d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2} UTC$")


def test_git_provenance_stamp_format_summary_report():
    stamp = _stamp_summary()
    assert _COMMIT_RE.match(stamp), stamp


def test_git_provenance_stamp_format_full_review():
    stamp = _stamp_full_review()
    assert _COMMIT_RE.match(stamp), stamp


def test_xlsx_glossary_sheet_has_generated_row(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    rows = [("TNA", None)]  # _row_to_record's no-candidates shape
    _write_xlsx("stamp_test", rows)

    out_path = tmp_path / "output" / "stamp_test.xlsx"
    assert out_path.exists()
    wb = load_workbook(out_path)
    def_ws = wb["Column Definitions"]
    values = [(r[0].value, r[1].value) for r in def_ws.iter_rows(min_row=2)]
    generated = dict(values).get("Generated")
    assert generated is not None
    assert _COMMIT_RE.match(generated), generated
    # every real column def is still present (stamp row is additive, not a replacement)
    assert set(COLUMN_DEFS.keys()).issubset({k for k, _ in values})


def test_csv_has_leading_generated_comment(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    rows = [("TNA", None)]
    _write_csv("stamp_test", rows)

    out_path = tmp_path / "output" / "stamp_test.csv"
    assert out_path.exists()
    first_line = out_path.read_text().splitlines()[0]
    assert first_line.startswith("# Generated: ")
    assert _COMMIT_RE.match(first_line[len("# Generated: "):]), first_line

    # the real header row is the *second* line, still parseable by a
    # comment-aware reader without special-casing the stamp
    second_line = out_path.read_text().splitlines()[1]
    assert second_line.split(",")[0] == list(COLUMN_DEFS.keys())[0]


def test_git_provenance_stamp_falls_back_to_unknown_on_git_failure(monkeypatch):
    """Best-effort contract: a git failure must never raise, and the stamp
    must clearly say so rather than silently omitting itself."""
    import scripts.candidate_summary_report as mod

    def _fake_state(files):
        return None, None

    monkeypatch.setattr("run_optimization_sweep._current_kernel_git_state", _fake_state)
    stamp = mod._git_provenance_stamp()
    assert stamp.startswith("commit=unknown at ")
